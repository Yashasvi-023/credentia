import re
from pptx import Presentation


PLACEHOLDER_PATTERN = r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}"


def extract_placeholders(file_path: str) -> list[str]:
    presentation = Presentation(file_path)

    placeholders = set()

    for slide in presentation.slides:
        for shape in slide.shapes:

            if not hasattr(shape, "text"):
                continue

            matches = re.findall(
                PLACEHOLDER_PATTERN,
                shape.text
            )

            placeholders.update(matches)

    return sorted(placeholders)