from pptx import Presentation



def replace_placeholders(
    template_path: str,
    output_path: str,
    recipient_data: dict
) -> None:

    presentation = Presentation(template_path)

    for slide in presentation.slides:
        for shape in slide.shapes:

            if not shape.has_text_frame:
                continue

            for paragraph in shape.text_frame.paragraphs:

                full_text = "".join(
                    run.text for run in paragraph.runs
                )

                updated_text = full_text

                for key, value in recipient_data.items():

                    placeholder_1 = "{{" + key + "}}"
                    placeholder_2 = "{{ " + key + " }}"

                    updated_text = updated_text.replace(
                        placeholder_1,
                        str(value)
                    )

                    updated_text = updated_text.replace(
                        placeholder_2,
                        str(value)
                    )

                if updated_text != full_text:

                    # Put the final text into the first run.
                    if paragraph.runs:
                        paragraph.runs[0].text = updated_text

                        # Clear the remaining runs.
                        for run in paragraph.runs[1:]:
                            run.text = ""

    presentation.save(output_path)